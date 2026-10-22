"""
One-off migration: wraps every not-yet-wrapped template formula in
IFERROR(...,"-"), so a formula with a missing/blank input shows a clean dash
instead of an Excel error - matching the convention already used by the
percentage/ratio columns and the newer Payout columns. This only edits
existing cell values in place (no column insert/delete, no merged-cell
changes), so it carries none of the corruption risk the column-rebuild
migration had.

Run once: python dev_scripts/wrap_formulas_iferror.py
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from app.services.template_store import get_template_bytes, save_template


def wrap_template(report_type: str) -> int:
    data = get_template_bytes(report_type)
    wb = load_workbook(io.BytesIO(data), data_only=False)
    ws = wb.active

    wrapped = 0
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(3, c)
        val = cell.value
        if isinstance(val, str) and val.startswith("=") and "IFERROR" not in val.upper():
            inner = val[1:]  # drop the leading "="
            cell.value = f'=IFERROR({inner},"-")'
            wrapped += 1

    if wrapped == 0:
        print(f"{report_type}: nothing to wrap")
        return 0

    out = io.BytesIO()
    wb.save(out)
    save_template(report_type, out.getvalue(), f"{report_type}.xlsx", "iferror-migration-script")
    print(f"{report_type}: wrapped {wrapped} formula(s)")
    return wrapped


if __name__ == "__main__":
    total = 0
    for rtype in ["deals", "pipeline_high", "pipeline_low", "leads", "meetings"]:
        total += wrap_template(rtype)
    print(f"\ntotal formulas wrapped: {total}")
